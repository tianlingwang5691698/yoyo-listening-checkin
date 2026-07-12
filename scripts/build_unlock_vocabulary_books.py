#!/usr/bin/env python3
import argparse
import json
import re
from collections import OrderedDict
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path("/Users/wangtianlong/工作/01_教学与备考/unlock 第二版")
SOURCES = {
    (1, "ls"): ROOT / "Unlock1/Unlock1 配套词汇/Unlock1_LS_wordlist.xlsx",
    (1, "rw"): ROOT / "Unlock1/Unlock1 配套词汇/Unlock1_RW_wordlist.xlsx",
    (2, "ls"): ROOT / "Unlock2/unlock2配套词汇/Unlock2_LS_wordlist.xlsx",
    (2, "rw"): ROOT / "Unlock2/unlock2配套词汇/Unlock2_RW_wordlist.xlsx",
    (3, "ls"): ROOT / "L3/Unlock3 配套词汇/Unlock3_LS_wordlist.xlsx",
    (3, "rw"): ROOT / "L3/Unlock3 配套词汇/Unlock3_RW_wordlist.xlsx",
    (4, "ls"): ROOT / "L4/Unlock4 配套词汇和quiz/LS/Vocab list",
    (4, "rw"): ROOT / "L4/Unlock4 配套词汇和quiz/RW/Vocab list",
}

POS_LABELS = {
    "n": "n.", "noun": "n.", "v": "v.", "verb": "v.",
    "adj": "adj.", "adjective": "adj.", "adv": "adv.", "adverb": "adv.",
    "prep": "prep.", "preposition": "prep.", "conj": "conj.", "conjunction": "conj.",
    "pron": "pron.", "pronoun": "pron.", "det": "det.", "determiner": "det.",
    "num": "num.", "numeral": "num.", "phr": "phr.", "phrase": "phr.",
    "nphr": "n. phr.", "nounphrase": "n. phr.", "vphr": "v. phr.",
    "verbphrase": "v. phr.", "phrasalverb": "v. phr.", "phrverb": "v. phr.",
    "adjphr": "adj. phr.", "adjectivephrase": "adj. phr.",
    "advphr": "adv. phr.", "adverbphrase": "adv. phr.",
    "prepphr": "prep. phr.", "prepositionalphrase": "prep. phr.",
    "conjphr": "conj. phr.", "conjunctionphrase": "conj. phr.",
}


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def clean_word(value):
    return clean(value).lstrip("*★•").strip()


def unit_number(value, fallback=0):
    match = re.search(r"(?:unit|u)\s*([1-8])", str(value), re.I)
    return int(match.group(1)) if match else fallback


def definition(pos, meaning):
    pos = clean(pos).strip("()")
    meaning = clean(meaning)
    return f"{pos} {meaning}".strip()


def normalize_pos(value):
    parts = clean(value).split("/")
    normalized = []
    for part in parts:
        key = re.sub(r"[.\s-]+", "", part.lower())
        if key not in POS_LABELS:
            return ""
        normalized.append(POS_LABELS[key])
    return "/".join(normalized)


def merge_definitions(definitions):
    groups = OrderedDict()
    loose = []
    for value in definitions:
        value = clean(value)
        match = re.search(r"[\u3400-\u9fff]", value)
        label = normalize_pos(value[:match.start()]) if match and match.start() > 0 else ""
        meaning = clean(value[match.start():]) if label else ""
        if not label or not meaning:
            if value and value not in loose:
                loose.append(value)
            continue
        groups.setdefault(label, [])
        if meaning not in groups[label]:
            groups[label].append(meaning)
    return [f"{label} {'；'.join(meanings)}" for label, meanings in groups.items()] + loose


def valid_word(value):
    value = clean_word(value)
    return bool(value and re.search(r"[A-Za-z]", value) and not re.search(r"单词|英文|vocabulary|wordlist|\b(?:unlock|ulk)\s*\d", value, re.I))


def read_xlsx(path):
    rows = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet_index, sheet in enumerate(workbook.worksheets, 1):
        unit = unit_number(sheet.title, sheet_index)
        blank_run = 0
        for values in sheet.iter_rows(values_only=True):
            values = list(values)
            if not any(clean(value) for value in values):
                blank_run += 1
                if blank_run >= 20:
                    break
                continue
            blank_run = 0
            if len(values) >= 5 and isinstance(values[0], (int, float)):
                word, phonetic, pos, meaning = values[1:5]
            elif len(values) >= 4 and valid_word(values[0]):
                word, phonetic, pos, meaning = values[:4]
            else:
                continue
            word = clean_word(word)
            if not valid_word(word) or not clean(meaning):
                continue
            rows.append({
                "word": word,
                "phonetic": clean(phonetic),
                "definitions": [definition(pos, meaning)],
                "example": "",
                "unit": unit,
                "sourceFile": path.name,
            })
    return rows


def split_l4_word_meaning(value):
    value = clean(value.replace("\n", " "))
    match = re.search(r"[\u3400-\u9fff]", value)
    if not match:
        return value, ""
    return value[:match.start()].strip(" -–—"), value[match.start():].strip()


def read_docx_dir(path):
    rows = []
    for docx_path in sorted(path.glob("*VocabList.docx"), key=lambda item: unit_number(item.name, 99)):
        unit = unit_number(docx_path.name)
        document = Document(docx_path)
        for table in document.tables:
            for row in table.rows[1:]:
                cells = [clean(cell.text) for cell in row.cells]
                if len(cells) < 4:
                    continue
                word, meaning = split_l4_word_meaning(cells[1])
                word = clean_word(word)
                if not valid_word(word) or not meaning:
                    continue
                rows.append({
                    "word": word,
                    "phonetic": cells[2],
                    "definitions": [definition(cells[3], meaning)],
                    "example": cells[5] if len(cells) > 5 else "",
                    "unit": unit,
                    "sourceFile": docx_path.name,
                })
    return rows


def merge_rows(rows, level, section):
    merged = OrderedDict()
    for row in rows:
        key = clean(row["word"]).lower()
        current = merged.get(key)
        if current is None:
            current = {
                "word": row["word"],
                "wordLower": key,
                "level": f"unlock{level}-{section}",
                "unlockLevel": level,
                "section": section.upper(),
                "units": [],
                "phonetic": row["phonetic"],
                "definitions": [],
                "example": row["example"],
                "source": "unlock-second-edition",
                "sourceFiles": [],
            }
            merged[key] = current
        if row["unit"] and row["unit"] not in current["units"]:
            current["units"].append(row["unit"])
        for item in row["definitions"]:
            if item and item not in current["definitions"]:
                current["definitions"].append(item)
        if not current["phonetic"] and row["phonetic"]:
            current["phonetic"] = row["phonetic"]
        if not current["example"] and row["example"]:
            current["example"] = row["example"]
        if row["sourceFile"] not in current["sourceFiles"]:
            current["sourceFiles"].append(row["sourceFile"])
    for current in merged.values():
        current["definitions"] = merge_definitions(current["definitions"])
    return list(merged.values())


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="data/unlock-vocabulary/output")
    args = parser.parse_args()
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    report = {"books": [], "total": 0}
    for (level, section), source in SOURCES.items():
        if not source.exists():
            raise FileNotFoundError(source)
        raw = read_xlsx(source) if source.suffix == ".xlsx" else read_docx_dir(source)
        rows = merge_rows(raw, level, section)
        target = output / f"unlock-{level}-{section}.json"
        target.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        unit_counts = {}
        for unit in range(1, 9):
            unit_rows = [row for row in rows if unit in row.get("units", [])]
            unit_target = output / f"unlock-{level}-u{unit}-{section}.json"
            unit_target.write_text(json.dumps(unit_rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            unit_counts[f"U{unit}"] = len(unit_rows)
        report["books"].append({
            "level": level,
            "section": section.upper(),
            "rawCount": len(raw),
            "count": len(rows),
            "unitCounts": unit_counts,
            "source": str(source),
            "output": str(target),
        })
        report["total"] += len(rows)
    (output / "clean-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
