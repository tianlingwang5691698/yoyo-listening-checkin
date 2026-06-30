#!/usr/bin/env python3
import json
import re
import subprocess
from pathlib import Path


ROOTS = [
    Path('/Users/wangtianlong/工作/未命名文件夹/3. 上海中考英语一模二模（12-24）/一模'),
    Path('/Users/wangtianlong/工作/未命名文件夹/6. 2025年上海一模'),
    Path('/Users/wangtianlong/工作/未命名文件夹/8.2026年上海一模'),
]
FORMAL_OUT = Path('data/imports/shanghai-em1-2014-2026/formal')
UPLOAD_OUT = Path('data/reading-em1')
DISTRICTS = ['黄浦', '黄埔', '徐汇', '长宁', '静安', '普陀', '虹口', '杨浦', '闵行', '宝山', '嘉定', '浦东', '浦东新区', '金山', '松江', '青浦', '奉贤', '崇明']
SECTION_META = {
    'A': {'sectionLabel': '阅读选择', 'difficultyLevel': 1, 'difficultyLabel': '基础理解'},
    'B': {'sectionLabel': '完形填空', 'difficultyLevel': 2, 'difficultyLabel': '语境词汇'},
}


def clean(text):
    text = str(text or '').replace('\f', ' ').replace('\u3000', ' ')
    return re.sub(r'\s+', ' ', text).strip()


def read_with_command(command):
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, timeout=90)
    return result.stdout


def read_text(path):
    suffix = path.suffix.lower()
    if suffix == '.doc':
        return read_with_command(['textutil', '-convert', 'txt', '-stdout', str(path)])
    if suffix == '.pdf':
        return read_with_command(['pdftotext', '-layout', str(path), '-'])
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if clean(p.text)]
        for table in doc.tables:
            for row in table.rows:
                cells = [clean(c.text) for c in row.cells if clean(c.text)]
                if cells:
                    parts.append(' '.join(cells))
        return '\n'.join(parts)
    return ''


def year_of(path):
    m = re.search(r'(20\d{2}|201\d)', str(path))
    return int(m.group(1)) if m else None


def district_of(path):
    text = str(path)
    for d in DISTRICTS:
        if d in text:
            return {'浦东新区': '浦东', '黄埔': '黄浦'}.get(d, d)
    return ''


def is_english_em1(path):
    s = str(path)
    return '英语' in s and '一模' in s and year_of(path) in range(2014, 2027) and district_of(path)


def is_paper_candidate(path):
    s = str(path)
    name = path.name
    if path.suffix.lower() not in {'.docx', '.doc', '.pdf'} or not is_english_em1(path):
        return False
    if any(x in s for x in ['语文', '数学', '物理', '化学', '道法', '历史', '跨学科', '答题卡', '听力文本', '听力文字']):
        return False
    if '参考答案' in name or ('答案' in name and not any(x in name for x in ['解析', '试题', '试卷'])):
        return False
    return True


def is_answer_candidate(path):
    s = str(path)
    name = path.name
    if path.suffix.lower() not in {'.docx', '.doc', '.pdf'} or not is_english_em1(path):
        return False
    if any(x in s for x in ['语文', '数学', '物理', '化学', '道法', '历史', '跨学科']):
        return False
    return any(x in name for x in ['答案', '解析', '参考'])


def file_score(path):
    s = str(path)
    score = 0
    if path.suffix.lower() == '.docx':
        score += 30
    if '解析版' in s or '解析' in s:
        score += 25
    if '试卷' in s or '试题' in s or '官方' in s:
        score += 12
    if '原卷' in s:
        score -= 6
    if path.suffix.lower() == '.pdf':
        score -= 12
    return score


def collect_files():
    papers = {}
    answers = {}
    for root in ROOTS:
        if not root.exists():
            continue
        for path in root.rglob('*'):
            if not path.is_file():
                continue
            key = (year_of(path), district_of(path))
            if not key[0] or not key[1]:
                continue
            if is_paper_candidate(path) and (key not in papers or file_score(path) > file_score(papers[key])):
                papers[key] = path
            if is_answer_candidate(path):
                answers.setdefault(key, []).append(path)
    return papers, answers


def answer_map(text):
    out = {}
    for start, end, answers in re.findall(r'(\d{1,2})\s*[-－—]\s*(\d{1,2})\s*([A-D]{2,})\b', text):
        start_num, end_num = int(start), int(end)
        if len(answers) == end_num - start_num + 1:
            for offset, answer in enumerate(answers):
                out[start_num + offset] = answer
    for num, ans in re.findall(r'(?:^|\s)(\d{1,2})\s*[\.．、]?\s*([A-D])\b', text):
        out[int(num)] = ans
    for num, ans in re.findall(r'【答案】\s*(\d{1,2})?\s*([A-D])\b', text):
        if num:
            out[int(num)] = ans
    return out


def between(text, start_re, end_re):
    start = re.search(start_re, text, re.I)
    if not start:
        return ''
    tail = text[start.start():]
    end = re.search(end_re, tail, re.I)
    return tail[:end.start()] if end else tail


def parse_options(chunk):
    opts = {}
    for m in re.finditer(r'(?:^|\s)([A-D])[\)）\.．、]\s*(.*?)(?=\s+[A-D][\)）\.．、]\s*|$)', clean(chunk), re.S):
        opts[m.group(1)] = clean(m.group(2))
    return opts


def invalid_composite_choice(prompt, options):
    if re.search(r'\b\d{1,2}[\.．]\s*_+', clean(prompt)):
        return True
    option_text = ''.join(clean(options.get(k)) for k in ['A', 'B', 'C', 'D'])
    return bool(re.fullmatch(r'[①②③④⑤⑥⑦⑧⑨⑩]+', option_text))


def parse_choice_questions(section, answers):
    text = clean(section)
    positions = []
    for m in re.finditer(r'(?:^|\s)(\d{1,2})[\.．]\s+', text):
        lookahead = text[m.end():m.end() + 360]
        if not all(re.search(rf'(?:^|\s){key}[\)）\.．、]\s*', lookahead) for key in ['A', 'B', 'C', 'D']):
            continue
        positions.append((int(m.group(1)), m.start(), m.end()))
    questions = []
    for idx, (number, start, content_start) in enumerate(positions):
        if number < 50 or number > 80:
            continue
        end = positions[idx + 1][1] if idx + 1 < len(positions) else len(text)
        chunk = text[content_start:end]
        a_pos = re.search(r'\s+A[\)）\.．、]\s*', chunk)
        if not a_pos:
            continue
        prompt = clean(chunk[:a_pos.start()])
        opts = parse_options(chunk[a_pos.start():])
        answer = answers.get(number, '')
        if not prompt or not all(opts.get(k) for k in ['A', 'B', 'C', 'D']) or answer not in ['A', 'B', 'C', 'D']:
            continue
        if invalid_composite_choice(prompt, opts):
            continue
        questions.append({
            'number': number,
            'prompt': prompt,
            'options': {k: opts[k] for k in ['A', 'B', 'C', 'D']},
            'answer': answer,
            'questionType': 'choice',
        })
    return questions


def passage_before_questions(section):
    text = clean(section)
    match = None
    for m in re.finditer(r'(?:^|\s)\d{1,2}[\.．]\s+', text):
        lookahead = text[m.end():m.end() + 360]
        if all(re.search(rf'(?:^|\s){key}[\)）\.．、]\s*', lookahead) for key in ['A', 'B', 'C', 'D']):
            match = m
            break
    passage = text[:match.start()] if match else text
    passage = re.sub(r'^(?:VI|Ⅵ)\.?\s*Reading comprehension.*?', '', passage, flags=re.I)
    passage = re.sub(r'^[AB]\.?\s*Choose the best.*?(?:\)|）)?', '', passage, flags=re.I)
    return clean(passage)


def section_block(text, section):
    if section == 'A':
        return between(text, r'A[\.．]\s*Choose the best answer', r'B[\.．]\s*Choose (?:the )?(?:best )?(?:answer|words|words or expressions).*?passage')
    return between(text, r'B[\.．]\s*Choose (?:the )?(?:best )?(?:answer|words|words or expressions).*?passage', r'C[\.．]\s*(?:Choose the words|Fill in the blanks|Read the passage)|D[\.．]\s*Answer|(?:VII|Ⅶ)\.?\s*Writing')


def build_item(year, district, source, section, passage, questions):
    meta = SECTION_META[section]
    return {
        '_id': f'sh-em1-{year}-{district}-reading-{section.lower()}',
        'title': f'{year} 上海{district}一模阅读 {section}',
        'year': year,
        'city': '上海',
        'district': district,
        'examType': '一模',
        'section': section,
        'sourceType': 'shanghai-mock',
        'sourceFile': source.name,
        'passage': passage,
        'questions': questions,
        'answerSentences': [],
        'phrases': [],
        'vocabulary': [],
        **meta,
    }


def main():
    papers, answer_files = collect_files()
    readings, report = [], []
    for (year, district), paper in sorted(papers.items()):
        try:
            paper_text = read_text(paper)
            answer_text = paper_text + '\n' + '\n'.join(read_text(p) for p in answer_files.get((year, district), [])[:4])
            answers = answer_map(answer_text)
            row = {'year': year, 'district': district, 'sourceFile': str(paper)}
            for section in ['A', 'B']:
                block = section_block(paper_text, section)
                passage = passage_before_questions(block)
                questions = parse_choice_questions(block, answers)
                row[f'reading{section}Questions'] = len(questions)
                row[f'reading{section}PassageChars'] = len(passage)
                if len(questions) >= 5 and len(passage) >= 300:
                    readings.append(build_item(year, district, paper, section, passage, questions))
            report.append(row)
        except Exception as exc:
            report.append({'year': year, 'district': district, 'sourceFile': str(paper), 'error': str(exc)})

    FORMAL_OUT.mkdir(parents=True, exist_ok=True)
    UPLOAD_OUT.mkdir(parents=True, exist_ok=True)
    text = json.dumps(readings, ensure_ascii=False, indent=2)
    (FORMAL_OUT / 'reading-passages.formal.json').write_text(text, encoding='utf-8')
    (FORMAL_OUT / 'reading-clean-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (UPLOAD_OUT / 'reading-passages.json').write_text(text, encoding='utf-8')
    (UPLOAD_OUT / 'reading-passages.js').write_text(f'module.exports = {text};\n', encoding='utf-8')
    (UPLOAD_OUT / 'README.md').write_text(
        '# 一模阅读上传目录\n\n上传 `reading-passages.json` 到云存储 `_content/reading-em1/reading-passages.json`。\n',
        encoding='utf-8',
    )
    print(json.dumps({
        'papers': len(papers),
        'readingPassages': len(readings),
        'outFile': str(UPLOAD_OUT / 'reading-passages.json'),
        'formalOut': str(FORMAL_OUT / 'reading-passages.formal.json'),
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
