#!/usr/bin/env python3
import json
import re
import subprocess
import urllib.request
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = Path("/Users/wangtianlong/工作/01_教学与备考/unlock 第二版")
REMOTE_BASE = "https://796f-youshengenglish-6glk12rd6c6e719b-1419984942.tcb.qcloud.la"
BACKUP_ROOT = ROOT / "data/unlock-vocabulary/backups/2026-07-13-before-examples"
OUTPUT_ROOT = ROOT / "data/unlock-vocabulary/examples-candidate"
TMP_ROOT = ROOT / "tmp/pdfs/unlock1"

LEVEL_ROOTS = {
    2: SOURCE_ROOT / "Unlock2/unlock2单词+单词Quiz",
    3: SOURCE_ROOT / "L3/unlock3配套单词和Quiz",
    4: SOURCE_ROOT / "L4/Unlock4 配套词汇和quiz",
}


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def word_key(value):
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def remote_path(level, unit, section):
    return f"dictionary_books/unlock-v2/level-{level}/unit-{unit}/{section}.json"


def local_path(root, level, unit, section):
    return root / f"level-{level}/unit-{unit}/{section}.json"


def download_formal_books():
    for level in range(1, 5):
        for unit in range(1, 9):
            for section in ("ls", "rw"):
                target = local_path(BACKUP_ROOT, level, unit, section)
                target.parent.mkdir(parents=True, exist_ok=True)
                url = f"{REMOTE_BASE}/{remote_path(level, unit, section)}?backup=20260713"
                with urllib.request.urlopen(urllib.request.Request(url, headers={"Cache-Control": "no-cache"}), timeout=30) as response:
                    payload = json.loads(response.read().decode("utf-8"))
                target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def source_examples():
    examples = {}
    for level, root in LEVEL_ROOTS.items():
        level_examples = {}
        for path in root.rglob("*.docx"):
            relative = str(path.relative_to(root)).lower()
            if "quiz" in relative or "答案" in path.name or "封面" in path.name:
                continue
            try:
                document = Document(path)
            except Exception:
                continue
            for table in document.tables:
                if not table.rows:
                    continue
                header = "|".join(clean(cell.text).lower() for cell in table.rows[0].cells)
                if "example" not in header:
                    continue
                for row in table.rows[1:]:
                    cells = [clean(cell.text) for cell in row.cells]
                    if len(cells) < 6 or not cells[5]:
                        continue
                    match = re.search(r"[\u3400-\u9fff]", cells[1])
                    word = (cells[1][:match.start()] if match else cells[1]).strip(" |,-–—")
                    if word:
                        level_examples.setdefault(word_key(word), cells[5])
        examples[level] = level_examples
    return examples


def recursive_strings(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, list):
        for item in value:
            yield from recursive_strings(item)
    elif isinstance(value, dict):
        for item in value.values():
            yield from recursive_strings(item)


def unlock1_corpus():
    TMP_ROOT.mkdir(parents=True, exist_ok=True)
    pdfs = [
        SOURCE_ROOT / "Unlock1/LS/Unlock 2e Listening and Speaking 1 Scripts.pdf",
        SOURCE_ROOT / "Unlock1/RW/Unlock 2e RW 1 Scripts.pdf",
        SOURCE_ROOT / "Unlock1/LS/【全彩】UNLOCK-1-听口.pdf",
        SOURCE_ROOT / "Unlock1/RW/【全彩】UNLOCK-1-读写.pdf",
    ]
    parts = []
    for index, pdf in enumerate(pdfs):
        target = TMP_ROOT / f"source-{index}.txt"
        subprocess.run(["pdftotext", str(pdf), str(target)], check=True)
        parts.append(target.read_text(encoding="utf-8", errors="ignore"))
    for path in (SOURCE_ROOT / "Unlock1/知识清单").rglob("*.docx"):
        try:
            document = Document(path)
        except Exception:
            continue
        parts.extend(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    for path in (ROOT / "data/transcript-build").rglob("*.json"):
        if "unlock1" not in str(path).lower():
            continue
        try:
            parts.extend(recursive_strings(json.loads(path.read_text(encoding="utf-8"))))
        except Exception:
            continue
    text = re.sub(r"\s+", " ", " ".join(parts))
    sentences = []
    blocked = re.compile(r"copyright|cambridge university press|photocopiable|isbn|www\.|unit \d+|track \d+", re.I)
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        sentence = clean(sentence).strip("0123456789 ")
        count = len(sentence.split())
        if 4 <= count <= 24 and re.search(r"[.!?]$", sentence) and re.search(r"[A-Za-z]", sentence) and not blocked.search(sentence):
            sentences.append(sentence)
    return sorted(set(sentences), key=lambda item: (len(item.split()), len(item)))


def corpus_example(word, sentences):
    variants = [clean(word)]
    if " - " in word:
        variants.extend(part.strip() for part in word.split(" - "))
    for variant in variants:
        pattern = re.compile(rf"(?<![A-Za-z]){re.escape(variant)}(?![A-Za-z])", re.I)
        for sentence in sentences:
            if pattern.search(sentence):
                return sentence
    return ""


def valid_example(example):
    example = clean(example)
    return bool(example and 4 <= len(example.split()) <= 28 and re.search(r"[.!?]$", example) and not re.search(r"[\u3400-\u9fff]", example))


def main():
    download_formal_books()
    docx_examples = source_examples()
    corpus = unlock1_corpus()
    missing = []
    report = {"books": [], "sourceExamples": 0, "corpusExamples": 0, "preservedExamples": 0, "missing": 0}
    for level in range(1, 5):
        for unit in range(1, 9):
            for section in ("ls", "rw"):
                source = local_path(BACKUP_ROOT, level, unit, section)
                rows = json.loads(source.read_text(encoding="utf-8"))
                book_stats = {"level": level, "unit": unit, "section": section, "total": len(rows), "filled": 0, "missing": 0}
                for index, row in enumerate(rows):
                    current = clean(row.get("example"))
                    example = current
                    origin = "preserved"
                    if level == 1:
                        example = corpus_example(row.get("word", ""), corpus)
                        origin = "unlock1-source"
                    elif level in (2, 3, 4) and not current:
                        example = docx_examples[level].get(word_key(row.get("word")), "")
                        origin = "companion-vocab-list"
                    if valid_example(example):
                        row["example"] = example
                        row["exampleSource"] = origin
                        book_stats["filled"] += 1
                        if origin == "unlock1-source":
                            report["corpusExamples"] += 1
                        elif origin == "companion-vocab-list":
                            report["sourceExamples"] += 1
                        else:
                            report["preservedExamples"] += 1
                    else:
                        row["example"] = ""
                        row.pop("exampleSource", None)
                        item_id = f"unlock-{level}-u{unit}-{section}-{index}"
                        missing.append({
                            "id": item_id,
                            "level": level,
                            "unit": unit,
                            "section": section.upper(),
                            "index": index,
                            "word": row.get("word", ""),
                            "definitions": row.get("definitions") or [],
                        })
                        book_stats["missing"] += 1
                        report["missing"] += 1
                target = local_path(OUTPUT_ROOT, level, unit, section)
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                report["books"].append(book_stats)
    (OUTPUT_ROOT / "missing-items.json").write_text(json.dumps(missing, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUTPUT_ROOT / "build-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
