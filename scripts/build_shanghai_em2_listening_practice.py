#!/usr/bin/env python3
import json
import os
import re
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from build_shanghai_em1_reading_upload import DISTRICTS, clean

DEFAULT_ROOTS = [
    Path('/Users/wangtianlong/工作/未命名文件夹/3. 上海中考英语一模二模（12-24）'),
    Path('/Users/wangtianlong/工作/未命名文件夹/7. 2025年上海二模/英语'),
    Path('/Users/wangtianlong/工作/未命名文件夹/9.2026年上海二模'),
]
ROOTS = [
    Path(p)
    for p in os.environ.get('SH_EM2_LISTENING_SOURCE_ROOTS', '').split(':')
    if p
] or DEFAULT_ROOTS
USE_IMPORT_TEXTS = not os.environ.get('SH_EM2_LISTENING_SOURCE_ROOTS')
OUT_DIR = Path('data/listening-em2')
IMAGE_DIR = OUT_DIR / 'images'
PRACTICE_OUT = OUT_DIR / 'listening-practice.json'
INDEX_OUT = Path('data/material-index.js')
INDEX_JSON_OUT = Path('data/material-index.json')
SOFFICE = Path('/Users/wangtianlong/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/soffice')


def is_em2_path(path):
    parts = [str(part) for part in Path(path).parts]
    if not any('二模' in part for part in parts):
        return False
    return not any(('一模' in part and '二模' not in part) for part in parts)


def year_of(path):
    m = re.search(r'(20\d{2}|201\d)', str(path))
    return int(m.group(1)) if m else None


def district_of(path):
    s = str(path)
    for d in DISTRICTS:
        if d in s:
            return {'浦东新区': '浦东', '黄埔': '黄浦'}.get(d, d)
    return ''


def display_year_of(source_year):
    return int(source_year) - 1 if source_year else source_year


def remap_audio_to_display_year(item, display_year):
    district = item.get('district')
    if not display_year or not district:
        return
    current_path = str(item.get('audioLocalPath') or '')
    if not current_path:
        return
    suffix = Path(current_path).suffix or '.mp3'
    target_name = f'sh-em2-{display_year}-{district}-listening{suffix}'
    target_path = OUT_DIR / 'audio' / target_name
    if not target_path.exists():
        return
    item['audioFile'] = target_name
    item['audioLocalPath'] = str(target_path)
    item['audioCloudPath'] = f'_content/listening-em2/audio/{target_name}'


def iter_docx_text(path):
    doc = Document(path)
    texts = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            text = clean(Paragraph(child, doc).text)
            if text:
                texts.append(text)
        elif isinstance(child, CT_Tbl):
            for row in Table(child, doc).rows:
                cells = [clean(cell.text) for cell in row.cells if clean(cell.text)]
                if cells:
                    texts.append(' '.join(cells))
    return '\n'.join(texts)


def read_with_command(args):
    proc = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    if proc.returncode != 0:
        return ''
    return proc.stdout.decode('utf-8', errors='ignore')


def read_text(path):
    if path.suffix.lower() == '.docx':
        try:
            return iter_docx_text(path)
        except Exception:
            return ''
    if path.suffix.lower() == '.doc':
        return read_with_command(['textutil', '-convert', 'txt', '-stdout', str(path)])
    if path.suffix.lower() == '.pdf':
        return read_with_command(['pdftotext', '-layout', str(path), '-'])
    if path.suffix.lower() == '.txt':
        return path.read_text(encoding='utf-8', errors='ignore')
    return ''


def candidate_text_files():
    files = []
    for root in ROOTS:
        if not root.exists():
            continue
        for path in root.rglob('*'):
            if path.name.startswith('~$') or not path.is_file():
                continue
            low = path.suffix.lower()
            s = str(path)
            if low in {'.doc', '.docx', '.pdf', '.txt'} and '英语' in s and is_em2_path(path):
                if any(bad in s for bad in ['答案纸', '答题纸', '答题卡']):
                    continue
                files.append(path)
    if USE_IMPORT_TEXTS:
        for path in Path('data/imports').glob('shanghai-em2-*/work/*.txt'):
            files.append(path)
    return sorted(files)


def score_text_source(path, text):
    score = 0
    s = str(path)
    if '原卷' in s or '真题卷' in s or '试卷' in s:
        score += 20
    if '解析' in s:
        score += 5
    if '听力文本' in s or '听力文字' in s or '参考答案' in s:
        score -= 20
    if 'Part 1' in text and 'Listening' in text:
        score += 20
    if re.search(r'6\.\s*A[).．]', text):
        score += 10
    if re.search(r'16\.\s*.+_{3,}', text):
        score += 10
    return score


def split_listening(text):
    start = re.search(r'Part\s*1\s+Listening|Listening Comprehension|Listening comprehension', text, re.I)
    if not start:
        return ''
    end = re.search(r'Part\s*2|Grammar and Vocabulary|Vocabulary and Grammar|第二部分', text[start.start():], re.I)
    if end:
        return text[start.start():start.start() + end.start()]
    return text[start.start():]


def option_map(line):
    opts = {}
    line = re.split(
        r'\s+(?=[CD][\).．]\s*Listen to (?:the dialogue|the passage|the conversation|the short passage|the recording))',
        clean(line),
        maxsplit=1,
        flags=re.I,
    )[0]
    for key, value in re.findall(r'([A-D])[\).．]\s*(.*?)(?=\s+[A-D][\).．]\s*|$)', clean(line)):
        value = clean(value)
        if value and not re.match(r'Listen (?:to (?:the dialogue|the passage|the conversation|the short passage|the recording)|and choose the right picture)', value, re.I):
            opts[key] = value
    return opts


def section_for_number(number, mode='20'):
    number = int(number)
    if mode == '25':
        if 1 <= number <= 6:
            return {
                'sectionKey': 'A',
                'sectionTitle': 'A. Listen and choose the right picture.'
            }
        if 7 <= number <= 14:
            return {
                'sectionKey': 'B',
                'sectionTitle': 'B. Listen and choose the best answer.'
            }
        if 15 <= number <= 20:
            return {
                'sectionKey': 'C',
                'sectionTitle': 'C. Listen and tell whether the statements are true or false.'
            }
        return {
            'sectionKey': 'D',
            'sectionTitle': 'D. Listen and complete the sentences.'
        }
    if 1 <= number <= 5:
        return {
            'sectionKey': 'A',
            'sectionTitle': 'A. Listen and choose the right picture.'
        }
    if 6 <= number <= 10:
        return {
            'sectionKey': 'B',
            'sectionTitle': 'B. Listen and choose the best answer.'
        }
    if 11 <= number <= 15:
        return {
            'sectionKey': 'C',
            'sectionTitle': 'C. Listen and tell whether the statements are true or false.'
        }
    return {
        'sectionKey': 'D',
        'sectionTitle': 'D. Listen and complete the sentences.'
    }


def apply_question_section(question):
    question.update(section_for_number(question.get('number', 0)))
    return question


def apply_question_sections(questions):
    mode = '25' if any(int(q.get('number') or 0) > 20 for q in questions) else '20'
    for question in questions:
        question.update(section_for_number(question.get('number', 0), mode))
    return questions


def parse_choice(lines):
    questions = []
    current = None
    for line in lines:
        m = re.match(r'^(\d{1,2})[\.．]\s*(.+)', line)
        if m and 6 <= int(m.group(1)) <= 14:
            if current:
                questions.append(current)
            num = int(m.group(1))
            current = {'number': num, 'prompt': 'Listen and choose the best answer.', 'questionType': 'choice', 'options': {}, 'answer': ''}
            opts = option_map(line)
            current['options'].update(opts)
            continue
        if current:
            opts = option_map(line)
            if opts:
                current['options'].update(opts)
    if current:
        questions.append(current)
    return [q for q in questions if len(q['options']) >= 4]


def parse_true_false(lines):
    questions = []
    for line in lines:
        m = re.match(r'^(1[1-9]|20)[\.．]?\s*(.+)', line)
        if m:
            if option_map(line):
                continue
            questions.append({
                'number': int(m.group(1)),
                'prompt': clean(m.group(2)),
                'questionType': 'truefalse',
                'options': {'T': 'T', 'F': 'F'},
                'answer': ''
            })
    return questions


def parse_blanks(lines):
    questions = []
    for line in lines:
        m = re.match(r'^(1[6-9]|2[0-5])[\.．]\s*(.+)', line)
        if m and (int(m.group(1)) >= 21 or re.search(r'_{3,}|__+', m.group(2))):
            questions.append({
                'number': int(m.group(1)),
                'prompt': clean(m.group(2)),
                'questionType': 'blank',
                'answer': ''
            })
    return questions


def extract_listening_answers(text):
    answers = {}
    candidates = []
    for marker in re.finditer(r'【答案】|参考答案|答案[:：]', text):
        block = text[marker.end():marker.end() + 3500]
        if re.search(r'1\s*-\s*5|1\s*-\s*6|6\s*-\s*10|7\s*-\s*14|11\s*-\s*15|15\s*-\s*20|16[\.．、]|1\s*[\.．、]\s*[A-H]', block, re.I):
            candidates.append(block)
    head = text[:2500]
    if re.search(r'Listening Comprehension|Listening comprehension|听力理解', head, re.I) and re.search(r'\b1\s*-\s*6\b|\b1\s*[\.．、]', head):
        candidates.insert(0, head)
    if not candidates:
        candidates = [text[:2500]]

    for raw_block in candidates:
        block = re.split(
            r'英语听力文字|听力文字|听力原文|Part\s*2|Vocabulary and Grammar|第二部分|II\.\s*(?:Choose|Vocabulary)',
            raw_block,
            maxsplit=1,
            flags=re.I,
        )[0]
        block = clean(block)

        for start, end, letters in re.findall(r'(\d{1,2})\s*-\s*(\d{1,2})\s*([A-HТTF](?:[\s\u00a0]*[A-HТTF])*)', block, flags=re.I):
            start = int(start)
            end = int(end)
            letters = re.sub(r'[^A-HТTF]', '', letters.upper()).replace('Т', 'T')
            if 1 <= start <= end <= 20 and end - start + 1 == len(letters):
                for offset, letter in enumerate(letters):
                    answers[start + offset] = letter

        for num, letter in re.findall(r'(?<!\d)(\d{1,2})\s*[\.．、]\s*\(([A-HТTF])\)', block, flags=re.I):
            num = int(num)
            if 1 <= num <= 20:
                answers[num] = letter.upper().replace('Т', 'T')
        for num, letter in re.findall(r'(?<!\d)(\d{1,2})\s*[\.．、]\s*([A-HТTF])(?=\s*\d{1,2}\s*[\.．、]|\s*$)', block, flags=re.I):
            num = int(num)
            if 1 <= num <= 20:
                answers[num] = letter.upper().replace('Т', 'T')
        for num, letter in re.findall(r'(?<!\d)(\d{1,2})\s*[\.．、]\s*([A-HТTF])(?=\s+[A-D][：:]|\s+Part|\s*$)', block, flags=re.I):
            num = int(num)
            if 1 <= num <= 20:
                answers[num] = letter.upper().replace('Т', 'T')

        blank_pattern = re.compile(
            r'(1[6-9]|2[0-5])\s*[\.．、]?\s*(.+?)(?=\s+(?:1[6-9]|2[0-5]|2[6-9]|3[0-9]|4[0-9]|5[0-9])\s*[\.．、]?|\s+\d{1,2}\s*-\s*\d{1,2}|\s+英语听力文字|\s+听力文字|\s+【听力原文】|\s+听力原文|\s+Part\s*(?:2|II)|\s+II\.|$)',
            re.I,
        )
        for num, value in blank_pattern.findall(block):
            num = int(num)
            if str(answers.get(num, '')).upper() in set('ABCDEFGHTF'):
                continue
            value = clean(value)
            value = re.sub(r'【解析】.*$', '', value).strip()
            value = re.sub(r'\s+Part\s*(?:2|II).*$|\s+II\..*$', '', value, flags=re.I).strip()
            if value and len(value) <= 80:
                answers[num] = value
        if is_valid_listening_answer_set(answers):
            break
    return answers


def extract_listening_transcript(text):
    if not text:
        return ''
    starts = []
    for marker in ['听力文稿', '听力文本', '听力文字', '录音稿', '录音文字']:
        pos = text.find(marker)
        if pos >= 0:
            starts.append(pos)
    q_pos = text.find('Q:')
    if q_pos >= 0:
        part_pos = text.rfind('Part 1 Listening', 0, q_pos)
        if part_pos >= 0:
            starts.append(part_pos)
        answer_pos = text.rfind('参考答案', 0, q_pos)
        if answer_pos >= 0:
            starts.append(answer_pos)
    if not starts:
        return ''
    source = text[max(starts):]
    end = re.search(r'Part\s*2|Grammar and Vocabulary|Vocabulary and Grammar|第二部分|II\.\s*Choose', source, re.I)
    if end:
        source = source[:end.start()]
    source = clean(source)
    if len(source) < 500:
        return ''
    if 'Q:' not in source and not re.search(r'\([A-F]\)', source):
        return ''
    return source[:12000]


def extract_docx_images(path, item_id):
    if path.suffix.lower() != '.docx':
        return []
    out = []
    try:
        doc = Document(path)
        rel_ids = []
        in_listening_image_section = False
        for paragraph in doc.paragraphs:
            text = clean(paragraph.text)
            if re.search(r'Listen and choose the right picture', text, re.I):
                in_listening_image_section = True
            elif in_listening_image_section and re.search(r'Part\s*2|Grammar and Vocabulary|Vocabulary and Grammar|第二部分', text, re.I):
                break
            if not in_listening_image_section:
                continue
            rel_ids.extend(paragraph._p.xpath('.//*[local-name()="blip"]/@*[local-name()="embed"]'))
        for rel_id in rel_ids[:12]:
            part = doc.part.related_parts.get(rel_id)
            if not part:
                continue
            ext = {
                'image/png': '.png',
                'image/jpeg': '.jpeg',
                'image/jpg': '.jpg',
                'image/gif': '.gif',
            }.get(part.content_type)
            if not ext:
                continue
            IMAGE_DIR.mkdir(parents=True, exist_ok=True)
            idx = len(out) + 1
            out_path = IMAGE_DIR / f'{item_id}-image-{idx}{ext}'
            out_path.write_bytes(part.blob)
            if out_path.stat().st_size <= 512:
                out_path.unlink()
                continue
            out.append({
                'localPath': str(out_path),
                'cloudPath': f'_content/listening-em2/images/{out_path.name}'
            })
    except Exception:
        return []
    return out


def convert_doc_to_docx(path, item_id):
    if path.suffix.lower() != '.doc' or not SOFFICE.exists():
        return None
    out_dir = IMAGE_DIR / '_converted-docx'
    out_dir.mkdir(parents=True, exist_ok=True)
    before = {p.name for p in out_dir.glob('*.docx')}
    subprocess.run(
        [str(SOFFICE), '--headless', '--convert-to', 'docx', '--outdir', str(out_dir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )
    converted = [p for p in out_dir.glob('*.docx') if p.name not in before]
    if not converted:
        converted = sorted(out_dir.glob(f'{path.stem}*.docx'), key=lambda p: p.stat().st_mtime, reverse=True)
    return converted[0] if converted else None


def extract_pdf_images(path, item_id):
    prefix = IMAGE_DIR / f'{item_id}-pdf'
    for stale in IMAGE_DIR.glob(f'{item_id}-pdf-*'):
        if stale.is_file():
            stale.unlink()
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    proc = subprocess.run(
        ['pdfimages', '-png', str(path), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )
    if proc.returncode != 0:
        return []
    out = []
    for idx, image_path in enumerate(sorted(IMAGE_DIR.glob(f'{item_id}-pdf-*')), 1):
        if idx > 12:
            image_path.unlink()
            continue
        ext = image_path.suffix.lower() or '.png'
        final_path = IMAGE_DIR / f'{item_id}-image-{idx}{ext}'
        if final_path.exists():
            final_path.unlink()
        image_path.rename(final_path)
        out.append({
            'localPath': str(final_path),
            'cloudPath': f'_content/listening-em2/images/{final_path.name}'
        })
    return out


def extract_images(path, item_id):
    if path.suffix.lower() == '.doc':
        converted = convert_doc_to_docx(path, item_id)
        return extract_docx_images(converted, item_id) if converted else []
    if path.suffix.lower() == '.docx':
        return extract_docx_images(path, item_id)
    if path.suffix.lower() == '.pdf':
        return extract_pdf_images(path, item_id)
    return []


def score_image_source(path, image_count):
    s = str(path)
    score = image_count * 10
    if '原卷' in s or '试卷' in s or '真题卷' in s:
        score += 30
    if '解析版' in s:
        score += 5
    if '听力文本' in s or '答案' in s or '参考答案' in s:
        score -= 40
    if path.suffix.lower() == '.docx':
        score += 5
    return score


def best_image_sources(keys):
    best = {}
    probe_dir = IMAGE_DIR / '_probe'
    probe_dir.mkdir(parents=True, exist_ok=True)
    for path in candidate_text_files():
        if path.suffix.lower() not in {'.doc', '.docx'}:
            continue
        year = year_of(path)
        district = district_of(path)
        key = (year, district)
        if key not in keys:
            continue
        probe_id = f'probe-{year}-{district}'
        old_image_dir = globals()['IMAGE_DIR']
        try:
            globals()['IMAGE_DIR'] = probe_dir
            images = extract_images(path, probe_id)
        finally:
            globals()['IMAGE_DIR'] = old_image_dir
        image_count = len(images)
        for image in images:
            p = Path(image['localPath'])
            if p.exists():
                p.unlink()
        if image_count < 1:
            continue
        score = score_image_source(path, image_count)
        old = best.get(key)
        if not old or score > old['score']:
            best[key] = {'path': path, 'imageCount': image_count, 'score': score}
    for stale in probe_dir.glob('*'):
        if stale.is_file():
            stale.unlink()
    return best


def parse_questions(text):
    sec = split_listening(text)
    if not sec:
        return []
    lines = [clean(x) for x in sec.splitlines() if clean(x)]
    questions = []
    questions.extend(parse_choice(lines))
    questions.extend(parse_true_false(lines))
    questions.extend(parse_blanks(lines))
    return sorted(apply_question_sections(questions), key=lambda x: x['number'])


def picture_questions(count=5):
    return [{
        'number': number,
        'prompt': 'Listen and choose the right picture.',
        'questionType': 'picture',
        'answer': ''
    } | section_for_number(number, '25' if count == 6 else '20') for number in range(1, count + 1)]


def best_sources():
    best = {}
    for path in candidate_text_files():
        year = year_of(path)
        district = district_of(path)
        if not year or not district:
            continue
        text = read_text(path)
        questions = parse_questions(text)
        if len(questions) < 10:
            continue
        key = (year, district)
        score = score_text_source(path, text) + len(questions)
        old = best.get(key)
        if not old or score > old['score']:
            best[key] = {'path': path, 'text': text, 'questions': questions, 'score': score}
    return best


def best_answer_sources():
    best = {}
    for path in candidate_text_files():
        year = year_of(path)
        district = district_of(path)
        if not year or not district:
            continue
        text = read_text(path)
        answers = extract_listening_answers(text)
        if not is_valid_listening_answer_set(answers):
            continue
        key = (year, district)
        s = str(path)
        score = len(answers) * 10
        if '解析版' in s or '答案' in s or '参考答案' in s:
            score += 50
        if '原卷版' in s:
            score -= 30
        old = best.get(key)
        if not old or score > old['score']:
            best[key] = {'path': path, 'answers': answers, 'score': score}
    return best


def best_transcript_sources():
    best = {}
    for path in candidate_text_files():
        year = year_of(path)
        district = district_of(path)
        if not year or not district:
            continue
        text = read_text(path)
        transcript = extract_listening_transcript(text)
        if not transcript:
            continue
        key = (year, district)
        score = len(transcript)
        if any(token in str(path) for token in ['听力文本', '听力原文', '听力文字', '听力文稿', '原文答案']):
            score += 5000
        old = best.get(key)
        if not old or score > old['score']:
            best[key] = {'path': path, 'transcript': transcript, 'score': score}
    return best


def is_valid_listening_answer_set(answers):
    if all(num in answers for num in range(1, 15)):
        tf_count = sum(1 for num in range(15, 21) if str(answers.get(num, '')).upper() in {'T', 'F'})
        blank_count = sum(
            1 for num in range(21, 26)
            if answers.get(num) and str(answers[num]).upper() not in set('ABCDEFGHTF')
        )
        return tf_count >= 4 and blank_count >= 3
    if not all(num in answers for num in range(1, 11)):
        return False
    tf_count = sum(1 for num in range(11, 16) if str(answers.get(num, '')).upper() in {'T', 'F'})
    blank_count = sum(
        1 for num in range(16, 21)
        if answers.get(num) and str(answers[num]).upper() not in set('ABCDEFGTF')
    )
    return tf_count >= 3 and blank_count >= 3


def main():
    sets = json.load(open(OUT_DIR / 'listening-sets.json', encoding='utf-8'))
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    for stale in IMAGE_DIR.glob('*'):
        if stale.is_file():
            stale.unlink()
    sources = best_sources()
    answer_sources = best_answer_sources()
    transcript_sources = best_transcript_sources()
    image_sources = best_image_sources(set(sources) | set(answer_sources))
    items = []
    for item in sets:
        key = (item.get('year'), item.get('district'))
        source = sources.get(key)
        if not source:
            continue
        next_item = dict(item)
        source_year = item.get('year')
        display_year = display_year_of(source_year)
        next_item['sourceYear'] = source_year
        next_item['year'] = display_year
        next_item['_id'] = f"sh-em2-{display_year}-{item.get('district')}-listening"
        next_item['title'] = f"{display_year} 上海{item.get('district')}二模听力"
        remap_audio_to_display_year(next_item, display_year)
        transcript_source = transcript_sources.get(key)
        if transcript_source:
            next_item['transcriptSourceFile'] = transcript_source['path'].name
            next_item['transcript'] = transcript_source['transcript']
            next_item['hasTranscript'] = True
        next_item['questions'] = source['questions']
        next_item['questionSourceFile'] = source['path'].name
        answer_source = answer_sources.get(key)
        if answer_source:
            next_item['answerSourceFile'] = answer_source['path'].name
            answers = answer_source['answers']
            for question in next_item['questions']:
                if question.get('number') in answers:
                    question['answer'] = answers[question['number']]
        image_source = image_sources.get(key)
        next_item['images'] = extract_images(image_source['path'], next_item['_id']) if image_source else []
        if image_source:
            next_item['imageSourceFile'] = image_source['path'].name
        next_item['hasPictureQuestions'] = bool(next_item['images'])
        if next_item['hasPictureQuestions']:
            existing = {q['number'] for q in next_item['questions']}
            picture_count = 6 if any(q.get('number') == 25 for q in next_item['questions']) else 5
            next_item['questions'] = picture_questions(picture_count) + [q for q in next_item['questions'] if q['number'] > picture_count]
            if answer_source:
                for question in next_item['questions']:
                    if question.get('number') in answers:
                        question['answer'] = answers[question['number']]
        if next_item['hasPictureQuestions'] and len(next_item['questions']) in {20, 25}:
            items.append(next_item)
    used_images = {Path(image['localPath']).name for item in items for image in item.get('images', [])}
    for image_path in IMAGE_DIR.glob('*'):
        if image_path.is_file() and image_path.name not in used_images:
            image_path.unlink()
    PRACTICE_OUT.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding='utf-8')

    material = json.load(open('data/material-index.json', encoding='utf-8'))
    slim = []
    for item in items:
        slim.append({
            '_id': item['_id'],
            'title': item['title'],
            'year': item['year'],
            'sourceYear': item.get('sourceYear'),
            'district': item['district'],
            'examType': item['examType'],
            'stage': item.get('stage', '初中'),
            'audioCloudPath': item.get('audioCloudPath', ''),
            'transcript': item.get('transcript', ''),
            'hasAudio': item.get('hasAudio', False),
            'hasTranscript': item.get('hasTranscript', False),
            'questionSourceFile': item.get('questionSourceFile', ''),
            'answerSourceFile': item.get('answerSourceFile', ''),
            'imageSourceFile': item.get('imageSourceFile', ''),
            'images': item.get('images', []),
            'questions': item.get('questions', [])
        })
    material['listeningEm2'] = slim
    INDEX_JSON_OUT.write_text(json.dumps(material, ensure_ascii=False, indent=2), encoding='utf-8')
    INDEX_OUT.write_text('module.exports = ' + json.dumps(material, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
    print(json.dumps({
        'practiceSets': len(items),
        'withImages': sum(1 for x in items if x.get('images')),
        'out': str(PRACTICE_OUT)
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
